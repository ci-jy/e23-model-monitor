"""Excel evidence pack (openpyxl).

Sheets:
* Summary      model-level alarm per slice, calibrated and rule-of-thumb
* Inventory    the inventory with validator flags, a live review-overdue
               formula and conditional formatting
* MON_<model>  per-metric results for every slice, followed by bin-count
               tables whose PSI/CSI is recomputed with live Excel formulas
               next to the Python value and their difference
"""
from __future__ import annotations

from datetime import date

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .inventory import COLUMNS, expected_rating

RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6EFCE")
HEAD = Font(bold=True)


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return None if np.isnan(v) else float(v)
    if v is pd.NaT:
        return None
    return v


def _header(ws, row: int, values: list[str]):
    for j, v in enumerate(values, 1):
        ws.cell(row, j, v).font = HEAD


def _autosize(ws, widths: dict[int, int] | None = None, default: int = 14):
    for j in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, default)


def write_inventory_sheet(wb: Workbook, inv: pd.DataFrame, violations: pd.DataFrame, as_of: date):
    ws = wb.create_sheet("Inventory")
    ws["A1"], ws["B1"] = "As of", as_of
    ws["A1"].font = HEAD
    ws["B1"].number_format = "yyyy-mm-dd"
    cols = COLUMNS + ["expected_rating", "flags", "review_overdue (formula)"]
    _header(ws, 3, cols)
    flags = violations.groupby("model_id")["rule"].apply(lambda s: "; ".join(s)).to_dict() if len(violations) else {}
    status_col = get_column_letter(COLUMNS.index("status") + 1)
    next_col = get_column_letter(COLUMNS.index("next_review") + 1)
    for i, (_, r) in enumerate(inv.iterrows(), start=4):
        for j, c in enumerate(COLUMNS, 1):
            cell = ws.cell(i, j, _cell_value(r.get(c, "")) if r.get(c, "") != "" else None)
            if isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
        ws.cell(i, len(COLUMNS) + 1, expected_rating(r["materiality"], r["complexity"], r["uses_ai_ml"], r["regulatory_use"]))
        ws.cell(i, len(COLUMNS) + 2, flags.get(r["model_id"], ""))
        ws.cell(i, len(COLUMNS) + 3,
                f'=IF(AND({status_col}{i}="Active",{next_col}{i}<>"",{next_col}{i}<$B$1),"OVERDUE","")')
    last = 3 + len(inv)
    flag_col = get_column_letter(len(COLUMNS) + 2)
    rng = f"A4:{get_column_letter(len(cols))}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${flag_col}4<>""'], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${flag_col}4=""'], fill=GREEN))
    ws.freeze_panes = "B4"
    _autosize(ws, {2: 40, len(COLUMNS) + 2: 50})


def _bin_table(ws, row: int, title: str, labels, ref_counts, cur_counts, smoothing: float, py_value: float,
               checks: list) -> int:
    """Write one bin table with live PSI formulas. Returns the next free row."""
    r = np.asarray(ref_counts, float)
    c = np.asarray(cur_counts, float)
    keep = (r + c) > 0
    ws.cell(row, 1, title).font = HEAD
    _header(ws, row + 1, ["bin", "reference count", "slice count", "reference share", "slice share", "PSI term"])
    first = row + 2
    labs = [l for l, k in zip(labels, keep) if k]
    last = first + len(labs) - 1
    for i, (lab, rc, cc) in enumerate(zip(labs, r[keep], c[keep])):
        k = first + i
        ws.cell(k, 1, lab)
        ws.cell(k, 2, float(rc))
        ws.cell(k, 3, float(cc))
        ws.cell(k, 4, f"=(B{k}+$B$3)/(SUM($B${first}:$B${last})+$B$3*ROWS($B${first}:$B${last}))")
        ws.cell(k, 5, f"=(C{k}+$B$3)/(SUM($C${first}:$C${last})+$B$3*ROWS($C${first}:$C${last}))")
        ws.cell(k, 6, f"=(E{k}-D{k})*LN(E{k}/D{k})")
        for col in (4, 5):
            ws.cell(k, col).number_format = "0.0000"
    t = last + 1
    ws.cell(t, 1, "PSI (Excel formula)").font = HEAD
    ws.cell(t, 6, f"=SUM(F{first}:F{last})")
    ws.cell(t + 1, 1, "PSI (Python)")
    ws.cell(t + 1, 6, float(py_value))
    ws.cell(t + 2, 1, "Difference")
    ws.cell(t + 2, 6, f"=F{t}-F{t + 1}")
    checks.append({"sheet": ws.title, "cell": f"F{t}", "python": float(py_value), "table": title})
    return t + 4


def write_model_sheet(wb: Workbook, model_id: str, ref, results: list[tuple[str, object, pd.DataFrame]],
                      checks: list):
    """``results``: list of (period, current Sample, evaluate_slice frame)."""
    ws = wb.create_sheet(f"MON_{model_id}"[:31])
    ws["A1"] = f"Monitoring results for {model_id}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Reference size"
    ws["B2"] = len(ref.sample)
    ws["A3"] = "Smoothing pseudo-count"
    ws["B3"] = ref.config.smoothing
    ws["C2"] = "Alpha (family-wise, Holm)"
    ws["D2"] = ref.config.alpha
    row = 5
    cols = ["period", "slice size", "metric", "value", "reference", "calibrated threshold", "p-value", "Holm p",
            "calibrated breach", "rule-of-thumb breach", "CI low", "CI high"]
    _header(ws, row, cols)
    row += 1
    for period, cur, df in results:
        for _, m in df.iterrows():
            vals = [period, len(cur), m["metric"], m["value"], m["reference"], m["calibrated_threshold"],
                    m["p_value"], m["p_holm"], bool(m["calibrated_breach"]), bool(m["rule_of_thumb_breach"]),
                    m.get("ci_low"), m.get("ci_high")]
            for j, v in enumerate(vals, 1):
                cell = ws.cell(row, j, _cell_value(v))
                if j in (4, 5, 6, 11, 12):
                    cell.number_format = "0.0000"
                if j in (7, 8):
                    cell.number_format = "0.00E+00"
            if m["calibrated_breach"]:
                ws.cell(row, 9).fill = RED
            if m["rule_of_thumb_breach"]:
                ws.cell(row, 10).fill = AMBER
            row += 1
    row += 2
    ws.cell(row, 1, "PSI / CSI recomputed from bin counts (live formulas)").font = Font(bold=True, size=12)
    row += 2
    for period, cur, df in results:
        for key, b in ref.binnings.items():
            col = cur.score if key == "score" else cur.features[key]
            name = "psi:score" if key == "score" else f"csi:{key}"
            py = float(df.loc[df.metric == name, "value"].iloc[0])
            row = _bin_table(ws, row, f"{period} {name}", b.labels(), ref.ref_counts[key], b.counts(col),
                             ref.config.smoothing, py, checks)
    _autosize(ws, {1: 26, 3: 16})


def write_summary_sheet(wb: Workbook, summary: pd.DataFrame):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "E-23 monitoring evidence pack"
    ws["A1"].font = Font(bold=True, size=13)
    cols = list(summary.columns)
    _header(ws, 3, cols)
    for i, (_, r) in enumerate(summary.iterrows(), start=4):
        for j, c in enumerate(cols, 1):
            ws.cell(i, j, _cell_value(r[c]))
        ca = cols.index("calibrated_alarm") + 1
        ws.cell(i, ca).fill = RED if r["calibrated_alarm"] else GREEN
        ra = cols.index("rule_of_thumb_alarm") + 1
        ws.cell(i, ra).fill = AMBER if r["rule_of_thumb_alarm"] else GREEN
    _autosize(ws, {cols.index("calibrated_breaches") + 1: 45, cols.index("rule_of_thumb_breaches") + 1: 45})


def build_workbook(summary: pd.DataFrame, model_results: dict, inventory=None, violations=None, as_of=None,
                   path="evidence_pack.xlsx") -> list[dict]:
    """Write the evidence pack; returns the list of formula PSI cells with the
    Python values they must reproduce."""
    wb = Workbook()
    wb.remove(wb.active)
    checks: list[dict] = []
    write_summary_sheet(wb, summary)
    if inventory is not None:
        write_inventory_sheet(wb, inventory, violations, as_of)
    for model_id, (ref, results) in model_results.items():
        write_model_sheet(wb, model_id, ref, results, checks)
    wb.save(path)
    return checks
