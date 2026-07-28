import subprocess
import sys

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook

from e23mon.cli import SAMPLE_DIR, main, run_monitoring
from e23mon.formula_check import recalculate, recalculate_and_compare
from e23mon.monitor import Config

@pytest.fixture(scope="module")
def sample_run(tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    cfg = yaml.safe_load((SAMPLE_DIR / "monitoring.yaml").read_text())
    res = run_monitoring(cfg, SAMPLE_DIR, out, Config(n_null=200, n_boot=50), SAMPLE_DIR / "inventory.csv",
                         verify=False)
    return res, out


def test_sample_outputs(sample_run):
    res, out = sample_run
    for f in ["breach_report.csv", "breach_summary.csv", "breach_report.md", "evidence_pack.xlsx",
              "inventory_violations.csv"]:
        assert (out / f).exists()
    s = res["summary"].set_index(["model_id", "period"])
    # Drifted months alarm on both models.
    for model in ["CR-PD-001", "CR-PD-002"]:
        for period in ["2026-07", "2026-08", "2026-09"]:
            assert s.loc[(model, period), "calibrated_alarm"]
    wb = load_workbook(out / "evidence_pack.xlsx")
    assert {"Summary", "Inventory", "MON_CR-PD-001", "MON_CR-PD-002"} <= set(wb.sheetnames)
    assert wb["Inventory"].conditional_formatting


def test_excel_formulas_match_python(sample_run, tmp_path):
    """The formula-computed PSI in every bin table matches Python to 1e-9,
    including tables with empty bins and an unseen category (July slice)."""
    res, out = sample_run
    check = recalculate_and_compare(out / "evidence_pack.xlsx", res["checks"])
    assert check["n_checked"] == len(res["checks"]) >= 100
    assert check["max_abs_diff"] < 1e-9
    titles = [c["table"] for c in res["checks"]]
    assert "2026-07 csi:EDUCATION" in titles


def test_inventory_overdue_formula(sample_run, tmp_path):
    res, out = sample_run
    wb = load_workbook(recalculate(out / "evidence_pack.xlsx", tmp_path), data_only=True)
    ws = wb["Inventory"]
    header = [c.value for c in ws[3]]
    idcol, fcol = header.index("model_id"), header.index("review_overdue (formula)")
    vals = {r[idcol].value: r[fcol].value for r in ws.iter_rows(min_row=4)}
    assert vals["PR-AUTO-004"] == "OVERDUE"
    assert not vals["CR-PD-001"]


def test_cli_validate(capsys):
    assert main(["validate", "examples/inventory.csv"]) == 0
    text = capsys.readouterr().out
    for rule in ["MISSING_OWNER", "REVIEW_OVERDUE", "DECOMMISSIONED_STILL_SCORING", "RATING_CONTRADICTS_CRITERIA"]:
        assert rule in text
    assert main(["validate", "examples/inventory.csv", "--strict"]) == 1


def test_cli_monitor_on_own_csvs(tmp_path):
    ref = pd.read_csv(SAMPLE_DIR / "reference.csv")
    cur = pd.read_csv(SAMPLE_DIR / "slice_2026-07.csv")
    ref.to_csv(tmp_path / "ref.csv", index=False)
    cur.to_csv(tmp_path / "jul.csv", index=False)
    out = tmp_path / "o"
    rc = subprocess.run([sys.executable, "-m", "e23mon", "monitor", "--reference", str(tmp_path / "ref.csv"),
                         "--slices", str(tmp_path / "jul.csv"), "--score-col", "pd_logit",
                         "--features", "LIMIT_BAL,EDUCATION,PAY_0", "--model-id", "MY-MODEL",
                         "--null-draws", "100", "--boot", "20", "--out", str(out)],
                        capture_output=True, text=True)
    assert rc.returncode == 0, rc.stderr
    rep = pd.read_csv(out / "breach_report.csv")
    assert set(rep.metric) == {"psi:score", "csi:LIMIT_BAL", "csi:EDUCATION", "csi:PAY_0", "auc_drop", "ks_drop",
                               "hl", "oe"}
    assert rep.model_id.eq("MY-MODEL").all()
